"""Excel / CSV 批量解析引擎 — 零依赖版 (csv 模块 + openpyxl)"""

import io
import csv
import os
import re
import subprocess
import tempfile
from typing import Optional, List, Dict, Any, Tuple

from services.query_engine import QueryEngine
from services.can_codec import CanCodec
from models.protocol import MessageDef


class ExcelParser:
    """CAN 抓包批量解析器 (无 pandas 版)"""

    def __init__(self, query: QueryEngine, codec: CanCodec):
        self.query = query
        self.codec = codec
        self._rows: List[Dict] = []
        self._columns: List[str] = []
        self._id_col: Optional[str] = None
        self._data_col: Optional[str] = None
        self._results: List[Dict] = []
        self.enable_tp = True
        self._cur_tp_key = None  # 当前活动 TP 组 key

    # ── 加载 ──────────────────────────────────────

    def load(self, filepath):
        """加载文件，支持路径 / bytes / BytesIO"""

        # bytes / BytesIO
        if isinstance(filepath, (bytes, io.BytesIO)):
            data = filepath if isinstance(filepath, bytes) else filepath.getvalue()
            try:
                self._try_xlsx(io.BytesIO(data))
                return
            except Exception:
                pass
            for method in [self._try_csv_text, self._try_fixed_width_text]:
                try:
                    if method(data, is_bytes=True):
                        return
                except Exception:
                    pass
            raise ValueError("无法识别文件格式")

        # 文件路径
        for method in [self._try_xlsx, self._try_csv_file, self._try_fixed_width_file]:
            try:
                if method(filepath):
                    return
            except Exception:
                continue

        # 都失败 → EFS 加密：用 cmd copy 解密到同目录后再试
        try:
            fd, dst = tempfile.mkstemp(
                suffix=".xlsx",
                dir=os.path.dirname(os.path.abspath(filepath))
            )
            os.close(fd)
            os.unlink(dst)
            r = subprocess.run(
                ["cmd", "/c", "copy", "/y", filepath, dst],
                capture_output=True, text=True, timeout=30
            )
            if r.returncode == 0 and os.path.exists(dst) and os.path.getsize(dst) > 0:
                for method in [self._try_xlsx, self._try_csv_file, self._try_fixed_width_file]:
                    try:
                        if method(dst):
                            os.unlink(dst)
                            return
                    except Exception:
                        pass
                os.unlink(dst)
        except Exception:
            pass

        raise ValueError("无法识别文件格式")

    def _try_xlsx(self, filepath):
        from openpyxl import load_workbook
        if isinstance(filepath, str):
            if not os.path.exists(filepath):
                raise FileNotFoundError
            wb = load_workbook(filepath, read_only=True, data_only=True)
        else:
            wb = load_workbook(filepath, read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        try:
            headers = [str(c) if c is not None else "" for c in next(it)]
        except StopIteration:
            wb.close()
            return False
        rows = [dict(zip(headers, row)) for row in it]
        wb.close()
        if rows:
            self._rows = rows
            self._columns = headers
            return True
        return False

    def _try_csv_file(self, filepath):
        with open(filepath, "rb") as f:
            data = f.read()
        return self._try_csv_text(data, is_bytes=True)

    def _try_csv_text(self, data, is_bytes=False):
        content = data
        for enc in ["utf-8-sig", "utf-8", "gbk"]:
            for sep in ["\t", ","]:
                try:
                    text = content.decode(enc)
                    reader = csv.reader(io.StringIO(text), delimiter=sep)
                    try:
                        headers = next(reader)
                    except StopIteration:
                        continue
                    rows = []
                    for row in reader:
                        if not any(row):
                            continue
                        d = {h: row[i] if i < len(row) else "" for i, h in enumerate(headers)}
                        if d:
                            rows.append(d)
                    if rows and len(headers) >= 2:
                        self._rows = rows
                        self._columns = headers
                        return True
                except Exception:
                    continue
        return False

    def _try_fixed_width_file(self, filepath):
        with open(filepath, "rb") as f:
            data = f.read()
        return self._try_fixed_width_text(data, is_bytes=True)

    def _try_fixed_width_text(self, data, is_bytes=False):
        content = data
        for enc in ["gbk", "utf-8", "utf-8-sig"]:
            try:
                text = content.decode(enc) if isinstance(content, bytes) else content
                lines = [l for l in text.splitlines() if l.strip()]
                if len(lines) < 2:
                    continue
                headers = [h.strip() for h in re.split(r'\s{2,}', lines[0].strip())]
                if len(headers) < 3:
                    continue
                rows = []
                for line in lines[1:]:
                    parts = [p.strip() for p in re.split(r'\s{2,}', line.strip())]
                    if len(parts) >= len(headers):
                        d = {h: parts[i] if i < len(parts) else "" for i, h in enumerate(headers)}
                        if d:
                            rows.append(d)
                    elif len(parts) >= len(headers) - 1:
                        d = {headers[i]: parts[i] if i < len(parts) - 1 else "" for i in range(len(headers) - 1)}
                        d[headers[-1]] = " ".join(parts[len(headers) - 1:])
                        rows.append(d)
                if rows and len(rows) > 1:
                    self._rows = rows
                    self._columns = headers
                    return True
            except Exception:
                continue
        return False

    # ── 列映射 ────────────────────────────────────

    def auto_map_columns(self) -> Tuple[Optional[str], Optional[str]]:
        if not self._columns:
            return None, None
        id_keywords = ["帧id", "canid", "can", "id", "标识符"]
        hex_keywords = ["hex", "数据(hex", "data(hex"]
        data_keywords = ["数据", "data", "报文", "dat"]
        id_col, data_col = None, None
        for col in self._columns:
            col_lower = str(col).lower().replace(" ", "").replace("_", "")
            if id_col is None:
                for kw in id_keywords:
                    if kw in col_lower:
                        id_col = col
                        break
        for col in self._columns:
            col_lower = str(col).lower().replace(" ", "").replace("_", "")
            for kw in hex_keywords:
                if kw in col_lower:
                    data_col = col
                    break
            if data_col:
                break
        if data_col is None:
            for col in self._columns:
                col_lower = str(col).lower().replace(" ", "").replace("_", "")
                for kw in data_keywords:
                    if kw in col_lower and "长度" not in str(col):
                        data_col = col
                        break
                if data_col:
                    break
        if data_col is None:
            for col in self._columns:
                col_lower = str(col).lower().replace(" ", "").replace("_", "")
                for kw in data_keywords:
                    if kw in col_lower:
                        data_col = col
                        break
                if data_col:
                    break
        self._id_col = id_col
        self._data_col = data_col
        return id_col, data_col

    def set_columns(self, id_col: str, data_col: str):
        self._id_col = id_col
        self._data_col = data_col

    @property
    def columns(self) -> List[str]:
        return self._columns

    @property
    def row_count(self) -> int:
        return len(self._rows)

    def preview(self, n: int = 20) -> List[Dict]:
        return self._rows[:n]

    # ── 解析 ──────────────────────────────────────

    def parse(self, progress_callback=None) -> List[Dict]:
        if not self._rows or self._id_col is None or self._data_col is None:
            raise ValueError("请先加载 Excel 并完成列映射")

        self._results = []
        total = len(self._rows)
        tp_groups: Dict[tuple, dict] = {}
        keep_cols = [c for c in self._columns
                     if c not in (self._id_col, self._data_col)
                     and not str(c).startswith("Unnamed")]

        for i, row in enumerate(self._rows):
            result = {"_index": i}
            for col in keep_cols:
                result[col] = row.get(col, "")

            try:
                can_id = self._parse_can_id(row.get(self._id_col, ""))
                result["CAN ID"] = f"0x{can_id:08X}"
                data = self._parse_data(row.get(self._data_col, ""))

                id_fields = self.query.parse_can_id(can_id)
                id_parts = []
                for sf in self.query.protocol.can_id_schema.fields:
                    v = id_fields.get(sf.name)
                    if v is not None and sf.fixed_value is None:
                        desc = sf.description or ""
                        id_parts.append(f"{sf.name}={v}(0x{v:X})" + (f"[{desc}]" if desc else ""))
                direction = self._get_direction(can_id, data, id_fields)
                result["CAN ID 解析"] = (direction + " | " if direction else "") + "; ".join(id_parts)
                result["数据域"] = data.hex(" ").upper() if data else ""

                pf = (can_id >> 16) & 0xFF

                if self.enable_tp and pf in (0xEB, 0xEC) and len(data) >= 5:
                    sa = can_id & 0xFF
                    ps = (can_id >> 8) & 0xFF
                    if pf == 0xEC and data[0] == 0x10:
                        # RTS 帧 → 开新组（每轮传输一个组）
                        key = (sa, ps, i)
                        tp_groups[key] = {"rts": {"can_id": can_id, "data": data},
                                          "dt_frames": [], "start_idx": i,
                                          "keep_cols": {c: row.get(c, "") for c in keep_cols}}
                        self._cur_tp_key = key
                    elif pf == 0xEB and self._cur_tp_key is not None:
                        # DT 帧挂到当前活动组（校验方向一致）
                        k = self._cur_tp_key
                        if (k[0], k[1]) == (sa, ps) and k in tp_groups:
                            seq = data[0]
                            tp_groups[k]["dt_frames"].append((seq, data))
                    continue

                msg = self.query.by_can_id(can_id, data)
                if msg is None:
                    result["报文"] = "UNKNOWN"
                    result["解析结果"] = ""
                else:
                    if msg.abbr.startswith("CMD_") or msg.abbr.startswith("RC_") or \
                       msg.abbr.startswith("GROUP") or msg.abbr.startswith("HEARTBEAT") or \
                       msg.abbr.startswith("PARAM") or msg.abbr.startswith("TELEMETRY"):
                        result["报文"] = f"{msg.abbr} [{msg.name}]"
                    else:
                        result["报文"] = msg.abbr
                    result["解析结果"] = self._format_decoded(msg, data)
            except Exception as e:
                result["报文"] = "ERROR"
                result["解析结果"] = str(e)[:80]

            self._results.append(result)
            if progress_callback:
                progress_callback(i + 1, total)

        # TP 重组
        for key, group in tp_groups.items():
            if not group["dt_frames"] or not group["rts"]:
                continue
            group["dt_frames"].sort(key=lambda x: x[0])
            frames_data = [f[1] for f in group["dt_frames"]]
            # 从 RTS 提取总长（标准 J1939: byte1-2 LE）
            rts_data = group["rts"]["data"]
            total_len = rts_data[1] | (rts_data[2] << 8)
            try:
                assembled = self.codec.tp_reassemble(frames_data, total_len=total_len)
            except Exception:
                continue
            if not assembled:
                continue

            target_msg = None
            target_pgn = rts_data[5] | (rts_data[6] << 8) | (rts_data[7] << 16)
            target_pf = target_pgn >> 8
            target_msg = self.query.by_pf(target_pf)

            if target_msg is None:
                for m in self.query.all_messages():
                    if m.tp and m.data_length and m.data_length == len(assembled):
                        target_msg = m
                        break

            if target_msg:
                result = {"_index": group["start_idx"], "_tp": True}
                result.update(group["keep_cols"])
                result["CAN ID"] = f"0x{group['rts']['can_id']:08X}" if group["rts"] else "TP"
                if group["rts"]:
                    id_fields = self.query.parse_can_id(group["rts"]["can_id"])
                    id_parts = [f"{sf.name}={id_fields.get(sf.name)}"
                                for sf in self.query.protocol.can_id_schema.fields
                                if id_fields.get(sf.name) is not None]
                    result["CAN ID 解析"] = "; ".join(id_parts)
                result["数据域"] = assembled.hex(" ").upper()
                result["报文"] = f"{target_msg.abbr} [TP]"
                result["解析结果"] = self._format_decoded(target_msg, assembled)
                self._results.append(result)

        self._results.sort(key=lambda r: r["_index"])
        return self._results

    def _format_decoded(self, msg, data: bytes) -> str:
        decoded = self.codec.decode(msg, data)
        parts = []
        for field in msg.fields:
            val = decoded.get(field.name, "")
            if isinstance(val, dict):
                val = val.get("label", str(val))
            if val == "" or val is None:
                continue
            unit = f"{field.unit}" if field.unit else ""
            parts.append(f"{field.name}={val}{unit}")
        return "; ".join(parts) if parts else ""

    def get_summary(self) -> Dict[str, int]:
        matched = sum(1 for r in self._results if r.get("报文") not in ("UNKNOWN", "ERROR", None))
        unknown = sum(1 for r in self._results if r.get("报文") == "UNKNOWN")
        error = sum(1 for r in self._results if r.get("报文") == "ERROR")
        tp = sum(1 for r in self._results if r.get("_tp"))
        return {"total": len(self._results), "matched": matched,
                "unknown": unknown, "error": error, "tp": tp}

    def to_rows(self) -> List[Dict]:
        return self._results

    # ── 导出 ──────────────────────────────────────

    def export(self, filepath: str):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        orig_cols = []
        for r in self._results[:1]:
            for k in r:
                if k not in ("_index", "_tp", "CAN ID", "数据域", "CAN ID 解析", "报文", "解析结果"):
                    orig_cols.append(k)
        display = orig_cols + ["CAN ID", "数据域", "CAN ID 解析", "报文", "解析结果"]
        ws.append(display)
        for r in self._results:
            ws.append([r.get(c, "") for c in display])
        wb.save(filepath)

    # ── 内部 ──────────────────────────────────────

    def _get_direction(self, can_id: int, data: bytes, id_fields: dict) -> str:
        matcher = self.query.protocol.message_matcher
        addrs = self.query.protocol.addresses
        if matcher.type == "data_field" and len(data) > 0:
            mt = data[0] & 0x0F
            return "监控→模块" if mt in (0x00, 0x02) else "模块→监控"
        if matcher.type == "command_number":
            dst = id_fields.get("dest_addr", 0)
            return "模块→监控" if dst in range(0xF0, 0xF9) else "监控→模块"
        sa = id_fields.get("sa", 0)
        ps = id_fields.get("ps", 0)
        sa_name = ps_name = ""
        for k, v in addrs.items():
            if v.get("value") == sa:
                sa_name = v.get("name", k)
            if v.get("value") == ps:
                ps_name = v.get("name", k)
        return f"{sa_name}→{ps_name}" if sa_name and ps_name else ""

    def _parse_can_id(self, value) -> int:
        if isinstance(value, (int, float)):
            return int(value)
        s = str(value).strip().lower().replace("0x", "").replace(" ", "")
        return int(s, 16)

    def _parse_data(self, value) -> bytes:
        if isinstance(value, bytes):
            return value
        s = str(value).strip().lower()
        for cm in ["//", "#", ";"]:
            idx = s.find(cm)
            if idx >= 0:
                s = s[:idx]
                break
        for prefix in ["0x", "(", "[", "'", '"']:
            if s.startswith(prefix):
                s = s[len(prefix):]
        for suffix in [")", "]", "'", '"']:
            if s.endswith(suffix):
                s = s[:-len(suffix)]
        s = ''.join(c for c in s if c in '0123456789abcdef ')
        s = s.replace(" ", "")
        if not s:
            return b''
        if len(s) % 2 != 0:
            s = '0' + s
        return bytes.fromhex(s)
